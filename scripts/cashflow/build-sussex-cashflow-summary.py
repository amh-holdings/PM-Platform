# One-tab Sussex cash flow for review, by SOV line and cost code.
#
# Phil, 2026-10-14: seven tabs were "too detailed for right now ... I only want
# one tab", then "can we have the cost codes and SOV so you can see where
# everything is". So: one sheet, months across -
#   CASH IN    one row per Exhibit E SOV line, then the total
#   CASH OUT   one row per cost code, then the total
#   NET CASH and CUMULATIVE CASH
# Amounts are the app's own, per source (from _export-sussex-cashflow-detail.mts,
# which checks the parts against the app's monthly totals); totals, net and
# cumulative are formulas.
#   python3 scripts/cashflow/build-sussex-cashflow-summary.py <detail.json> <out.xlsx>
import json, sys
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.properties import CalcProperties

d = json.load(open(sys.argv[1]))
OUT = sys.argv[2]
F, NAVY = "Arial", "1C2B4A"
font = lambda **k: Font(name=F, size=k.pop("size", 10), **k)
MONEY = '$#,##0;($#,##0);"-"'
fill = lambda c: PatternFill("solid", fgColor=c)
line = Border(top=Side(style="thin", color="333333"))

months = [m for m in d["app"] if abs(m["cashIn"]) > 0.005 or abs(m["cashOut"]) > 0.005]
M0 = 3                       # first month column (C)
MT = M0 + len(months)        # total column
ML = L(M0), L(MT - 1)

wb = Workbook()
ws = wb.active
ws.title = "Cash Flow"
p = d["project"]
ws["A1"] = f'{p["name"]} - Cash Flow by SOV Line and Cost Code'
ws["A1"].font = font(size=14, bold=True, color=NAVY)
ws["A2"] = f'IBT TVIG CT JV, LLC  |  Owner: {p["client"]}  |  As of {date.today().strftime("%b %-d, %Y")}  |  Months show when cash moves, after payment terms and retainage'
ws["A2"].font = font(italic=True, color="555555")

H = 4
for i, h in enumerate(["Item", "Description"] + [m["label"] for m in months] + ["Total"], 1):
    c = ws.cell(H, i, h)
    c.font, c.fill = font(bold=True, color="FFFFFF"), fill(NAVY)
    c.alignment = Alignment(horizontal="center")


def section(start, title, side, key):
    ws.cell(start, 1, title).font = font(bold=True, color=NAVY, size=11)
    r = start + 1
    srcs = [s for s in d["sources"] if s["side"] == side and any(abs(x[key]) > 0.005 for x in s["rows"])]
    for s in srcs:
        ws.cell(r, 1, s["ref"]).font = font()
        ws.cell(r, 2, s["label"]).font = font()
        by = {x["month"]: x[key] for x in s["rows"]}
        for i, m in enumerate(months):
            v = by.get(m["month"], 0)
            c = ws.cell(r, M0 + i, round(v, 2) if abs(v) > 0.005 else None)
            c.font, c.number_format = font(color="0000FF"), MONEY
        c = ws.cell(r, MT, f"=SUM({ML[0]}{r}:{ML[1]}{r})")
        c.font, c.number_format = font(bold=True), MONEY
        r += 1
    tot = r
    ws.cell(tot, 2, f"Total {title.title()}").font = font(bold=True)
    for col in range(M0, MT + 1):
        c = ws.cell(tot, col, f"=SUM({L(col)}{start+1}:{L(col)}{tot-1})")
        c.font, c.number_format = font(bold=True), MONEY
    for col in range(1, MT + 1):
        ws.cell(tot, col).fill = fill("E8ECF3")
        ws.cell(tot, col).border = line
    return tot


tin = section(H + 1, "CASH IN", "in", "cashIn")
tout = section(tin + 2, "CASH OUT", "out", "cashOut")

net = tout + 2
cum = net + 1
ws.cell(net, 2, "Net Cash").font = font(bold=True, size=11)
ws.cell(cum, 2, "Cumulative Cash").font = font(bold=True, size=11)
for i in range(len(months) + 1):
    col = L(M0 + i)
    c = ws[f"{col}{net}"]
    c.value = f"={col}{tin}-{col}{tout}"
    c.font, c.number_format = font(bold=True), MONEY
    if i < len(months):
        c = ws[f"{col}{cum}"]
        c.value = f"={col}{net}" if i == 0 else f"={L(M0 + i - 1)}{cum}+{col}{net}"
        c.font, c.number_format = font(bold=True), MONEY
for r in (net, cum):
    for col in range(1, MT + 1):
        ws.cell(r, col).fill = fill("D6DEEB")
ws.cell(net, 1).border = line

lo = cum + 2
ws.cell(lo, 2, "Lowest cumulative cash").font = font(bold=True)
c = ws.cell(lo, 3, f"=MIN({ML[0]}{cum}:{ML[1]}{cum})")
c.font, c.number_format, c.fill = font(bold=True), MONEY, fill("FFFF00")
ws.cell(lo, 4, f"=INDEX({ML[0]}{H}:{ML[1]}{H},MATCH(C{lo},{ML[0]}{cum}:{ML[1]}{cum},0))").font = font(bold=True)
ws.cell(lo + 1, 2, "Net cash at completion").font = font(bold=True)
c = ws.cell(lo + 1, 3, f"={L(MT)}{net}")
c.font, c.number_format = font(bold=True), MONEY
ws.cell(lo + 3, 1, "Cash In: Dimension payments on Exhibit E, Net 30, less 5% retainage (LNTP exempt); retainage returns 30 days after Substantial Completion.").font = font(italic=True, color="555555")
ws.cell(lo + 4, 1, "Cash Out: by cost code, after subcontractor / vendor terms and 10% sub retainage (paid at Substantial Completion). Blue = amounts from the PM Platform.").font = font(italic=True, color="555555")

ws.column_dimensions["A"].width = 10
ws.column_dimensions["B"].width = 42
for i in range(len(months)):
    ws.column_dimensions[L(M0 + i)].width = 11
ws.column_dimensions[L(MT)].width = 13
ws.freeze_panes = f"{L(M0)}{H+1}"
ws.sheet_view.showGridLines = False
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
wb.calculation = CalcProperties(fullCalcOnLoad=True)
wb.save(OUT)
print("saved", OUT, len(months), "months", "rows in", tin - H - 2, "out", tout - tin - 3)

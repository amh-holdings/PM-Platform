import json, sys
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.comments import Comment

d = json.load(open(sys.argv[1]))
OUT = sys.argv[2]
AS_OF = date.today().strftime("%b %-d, %Y")

F = "Arial"
NAVY = "1C2B4A"
f_title = Font(name=F, size=14, bold=True, color=NAVY)
f_sub = Font(name=F, size=10, italic=True, color="555555")
f_hdr = Font(name=F, size=10, bold=True, color="FFFFFF")
f_b = Font(name=F, size=10, bold=True)
f_n = Font(name=F, size=10)
f_in = Font(name=F, size=10, color="0000FF")      # hardcoded inputs
f_link = Font(name=F, size=10, color="008000")    # links to other sheets
fill_hdr = PatternFill("solid", fgColor=NAVY)
fill_tot = PatternFill("solid", fgColor="E8ECF3")
fill_key = PatternFill("solid", fgColor="FFFF00")
thin = Side(style="thin", color="BBBBBB")
top = Border(top=Side(style="thin", color="333333"))
MONEY = '$#,##0;($#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'

# Months that carry anything, Jun 2026 on.
months = [a for a in d["app"] if a["month"] >= "2026-06-01"]
active = [m for m in months if any(abs(m[k]) > 0.005 for k in ("billed", "cost", "cashIn", "cashOut"))]
first, last = active[0]["month"], active[-1]["month"]
months = [m for m in months if first <= m["month"] <= last]
MI = {m["month"]: i for i, m in enumerate(months)}

def hdr(ws, row, cols, widths=None):
    for i, h in enumerate(cols, 1):
        c = ws.cell(row=row, column=i, value=h)
        c.font, c.fill = f_hdr, fill_hdr
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[L(i)].width = w

wb = Workbook()

# ---------------------------------------------------------------- detail grids
def grid(name, side, key, title, note, basis_label, check=False):
    ws = wb.create_sheet(name)
    ws["A1"] = title; ws["A1"].font = f_title
    ws["A2"] = note; ws["A2"].font = f_sub
    fixed = ["Item", "Description", "Value ($)", basis_label]
    cols = fixed + [m["label"] for m in months] + ["Total ($)"] + (["Check vs value"] if check else [])
    hdr(ws, 4, cols, [10, 38, 13, 46] + [11] * len(months) + [13, 13])
    r = 5
    srcs = [s for s in d["sources"] if s["side"] == side and any(abs(x[key]) > 0.005 for x in s["rows"])]
    for s in srcs:
        ws.cell(r, 1, s["ref"]).font = f_n
        ws.cell(r, 2, s["label"]).font = f_n
        v = ws.cell(r, 3, round(s["value"], 2)); v.font = f_in; v.number_format = MONEY
        basis = s.get("terms") or f'{s["method"].title()} on {s["trigger"]}' + (f' | {s["retainage"]:.0%} retainage' if s.get("retainage") else " | no retainage")
        ws.cell(r, 4, basis).font = f_n
        by = {x["month"]: x[key] for x in s["rows"]}
        for m in months:
            amt = by.get(m["month"], 0)
            c = ws.cell(r, 5 + MI[m["month"]], round(amt, 2) if abs(amt) > 0.005 else None)
            c.font, c.number_format = f_in, MONEY
        tc = 5 + len(months)
        t = ws.cell(r, tc, f"=SUM({L(5)}{r}:{L(tc - 1)}{r})"); t.font, t.number_format = f_b, MONEY
        if check:
            k = ws.cell(r, tc + 1, f"={L(tc)}{r}-C{r}"); k.font, k.number_format = f_n, MONEY
        r += 1
    last_row = r - 1
    ws.cell(r, 2, "Total").font = f_b
    for c in [3] + list(range(5, 6 + len(months))):
        x = ws.cell(r, c, f"=SUM({L(c)}5:{L(c)}{last_row})")
        x.font, x.number_format, x.fill, x.border = f_b, MONEY, fill_tot, top
    ws.cell(r, 2).fill = fill_tot
    ws.freeze_panes = "E5"
    return ws, r

cin, cin_tot = grid("Cash In", "in", "cashIn", "Cash in from Dimension, by SOV line",
    "Month the money reaches the bank: billing month + Net 30, less 5% retainage, retainage returned 30 days after Substantial Completion. Blue = values from the PM Platform.",
    "How it bills")
bil, bil_tot = grid("Billing", "in", "billed", "Billed to Dimension, by SOV line",
    "Month each line is earned and invoiced (Exhibit E). Milestone lines bill when their schedule task finishes; install lines bill monthly as work is completed.",
    "How it bills", check=True)
cout, cout_tot = grid("Cash Out", "out", "cashOut", "Cash out, by cost code",
    "Month the money leaves the bank, after vendor / sub terms and retainage. Sub retainage is paid at Substantial Completion.",
    "Payment terms")
cost, cost_tot = grid("Cost", "out", "cost", "Cost incurred, by cost code",
    "Month the cost is incurred (work done / milestone reached), before payment terms.",
    "Payment terms")

# ---------------------------------------------------------------- summary
ws = wb.active; ws.title = "Summary"
p = d["project"]
ws["A1"] = f'{p["name"]} - Project Cash Flow'; ws["A1"].font = f_title
ws["A2"] = f'Owner: {p["client"]}  |  Contractor: IBT TVIG CT JV, LLC  |  As of {AS_OF}  |  Source: AHC PM Platform'; ws["A2"].font = f_sub

kp = [("Contract value", f"=SUM('Billing'!C5:C{bil_tot-1})", MONEY),
      ("Cost at completion", f"=SUM('Cost'!{L(5)}{cost_tot}:{L(4+len(months))}{cost_tot})", MONEY),
      ("Margin at completion", "=B5-B6", MONEY),
      ("Margin %", "=IF(B5=0,0,B7/B5)", PCT),
      ("Lowest cumulative cash", None, MONEY),
      ("Month of low point", None, None),
      ("Cash at completion", None, MONEY)]
for i, (lab, f, nf) in enumerate(kp):
    r = 5 + i
    ws.cell(r, 1, lab).font = f_b
    c = ws.cell(r, 2, f); c.font = f_link if f and "!" in f else f_n
    if nf: c.number_format = nf
    c.fill = fill_tot
ws.column_dimensions["A"].width = 26
for col, w in zip("BCDEFGHI", [16, 16, 16, 16, 16, 18, 16, 12]):
    ws.column_dimensions[col].width = w

H = 13
hdr(ws, H, ["Month", "Billed to Dimension", "Cost incurred", "Cash in", "Cash out", "Net cash", "Cumulative cash", "PM Platform cumulative", "Difference"])
ws.column_dimensions["A"].width = 26
for i, m in enumerate(months):
    r = H + 1 + i
    col = L(5 + i)
    ws.cell(r, 1, m["label"]).font = f_n
    for c, (sheet, tot) in zip((2, 3, 4, 5), (("Billing", bil_tot), ("Cost", cost_tot), ("Cash In", cin_tot), ("Cash Out", cout_tot))):
        x = ws.cell(r, c, f"='{sheet}'!{col}{tot}"); x.font, x.number_format = f_link, MONEY
    x = ws.cell(r, 6, f"=D{r}-E{r}"); x.font, x.number_format = f_n, MONEY
    x = ws.cell(r, 7, f"=F{r}" if i == 0 else f"=G{r-1}+F{r}"); x.font, x.number_format = f_b, MONEY
    x = ws.cell(r, 8, round(m["cumulative"], 2)); x.font, x.number_format = f_in, MONEY
    x = ws.cell(r, 9, f"=G{r}-H{r}"); x.font, x.number_format = f_n, MONEY
end = H + len(months)
tr = end + 1
ws.cell(tr, 1, "Total").font = f_b
for c in range(2, 7):
    x = ws.cell(tr, c, f"=SUM({L(c)}{H+1}:{L(c)}{end})"); x.font, x.number_format, x.fill, x.border = f_b, MONEY, fill_tot, top
ws.cell(tr, 1).fill = fill_tot
ws["B9"] = f"=MIN(G{H+1}:G{end})"
ws["B10"] = f"=INDEX(A{H+1}:A{end},MATCH(B9,G{H+1}:G{end},0))"
ws["B11"] = f"=G{end}"
for a in ("B9", "B11"): ws[a].number_format = MONEY
ws["B9"].fill = fill_key
ws.cell(tr + 2, 1, "Column H is the PM Platform's own cumulative cash for each month; column I should read zero.").font = f_sub
ws.cell(tr + 3, 1, "Blue = values from the PM Platform  |  Black = formulas  |  Green = links to the detail tabs  |  Yellow = the number to watch").font = f_sub
ws.freeze_panes = f"B{H+1}"

# ---------------------------------------------------------------- assumptions
wa = wb.create_sheet("Assumptions", 1)
wa["A1"] = "Assumptions"; wa["A1"].font = f_title
wa["A2"] = "Terms the cash flow is built on. Source: Phil Horwitch unless noted; contract documents where cited."; wa["A2"].font = f_sub
hdr(wa, 4, ["#", "Area", "Assumption", "Source"], [5, 22, 90, 34])
pa = d["payApps"]
A = [
 ("Owner billing", "Dimension pays Net 30 from each billing month.", "Exhibit E / contract"),
 ("Owner billing", "5% retainage on every billing except the LNTP lines (1.01.1-1.01.9), returned at Substantial Completion and paid Net 30 after it.", "Exhibit E Note 3"),
 ("Owner billing", "Milestone lines (LNTP, NTP, engineering, permits, equipment, MC / COD / SC) bill in full when their schedule task finishes.", "Exhibit E"),
 ("Owner billing", "Install lines 6.01-8.03 bill monthly as work is completed, spread over the days of their schedule task. Dimension allows monthly percent-complete billing.", "Phil / Dimension"),
 ("Owner billing", "Engineering plan sets bill per the LNTP when each package task completes.", "LNTP"),
 ("Owner billing", "Equipment can be billed at 50% when the PO is issued. Not yet modelled - switchgear will be added once its PO is loaded; today 5.06 bills on delivery.", "Phil"),
 ("Owner billing", f"AFP 1 (invoice 202-01, Jun 26 2026, $29,038.82 LNTP execution) is open and unpaid; it is expected in the current month. It is the only pay app issued.", "Invoice 202-01"),
 ("Owner billing", "4.00 Permits bills when all contractor permits are received (schedule 2.2).", "Phil"),
 ("Subcontractors", "Every subcontract carries 10% retainage and Net 30 (Net 45 will be sought). Until subs are signed, construction codes 300-xxxx and 200-1050 carry these as placeholders.", "Phil"),
 ("Subcontractors", "Pure Power Engineering (electrical design): Net 45, 10% retainage, paid on its SOV milestones. Exactus Energy (civil design): Net 30, no retainage.", "Subcontracts"),
 ("Subcontractors", "AC/MV cable (200-1050) is procured by the electrical subcontractor, paid as sub material at Electrical start.", "Phil"),
 ("Procurement", "Racking (ATI / Ampacity, 200-1030): contract signed at schedule row 27 (Structural Subcontract Executed, Oct 14 2026); 25% deposit Net 5 from signing; 75% on delivery Net 30.", "Racking contract"),
 ("Procurement", "Other vendor POs: 30% deposit Net 30 from PO, 70% on delivery Net 30.", "Phil"),
 ("Procurement", "Equipment order-by dates come from vendor lead times counted back from the install start (schedule section 3). Modules and inverters are owner-furnished.", "Vendor quotes / Phil"),
 ("General conditions", "TVIG Annex A (EPCA, $774,333) is paid exactly per the AHC Annex A cash flow: each month billed one month in advance, Dec 2026 - Nov 2027.", "Sussex_CSG1_Annex_A_Cash_Flow.xlsx"),
 ("General conditions", "IBT Annex A (EPCA, $95,000) is spread evenly over construction.", "Phil"),
 ("General conditions", "Bonds ($120,000) paid in full Nov 2026. Insurance ($100,000) paid one month before mobilization.", "Phil"),
 ("General conditions", "Other site general conditions (field office, travel, tools, utilities, safety, contingency) are spread over construction and paid the month incurred.", "Job cost budget"),
 ("Schedule", "Dates follow the PM Platform schedule as of today: mobilization after the Land Disturbance Permit; racking after the Building Permit; electrical after the Electrical Permit; IFC after all contractor permits; MC after Electrical; COD and SC after Commissioning.", "PM Platform schedule"),
 ("Margin", "Electrical design budget ($99,100) exceeds the Pure Power subcontract ($96,100); the $3,000 difference is counted as margin.", "Phil"),
]
for i, (area, text, src) in enumerate(A, 1):
    r = 4 + i
    wa.cell(r, 1, i).font = f_n
    wa.cell(r, 2, area).font = f_b
    c = wa.cell(r, 3, text); c.font = f_n; c.alignment = Alignment(wrap_text=True, vertical="top")
    wa.cell(r, 4, src).font = f_n
    wa.row_dimensions[r].height = 30

# ---------------------------------------------------------------- open items
wo = wb.create_sheet("Open Items", 2)
wo["A1"] = "Open items"; wo["A1"].font = f_title
wo["A2"] = "What is still a placeholder or unconfirmed, and how it would move the cash flow."; wo["A2"].font = f_sub
hdr(wo, 4, ["#", "Item", "Current treatment", "Effect when resolved"], [5, 34, 60, 60])
O = [
 ("October 2026 low point", "Racking deposit ($225K) goes out ~Oct 19; Dimension's 5.04 racking deposit ($358K) is expected ~Nov 13 on Net 30.", "Invoicing 5.04 at signing and getting it paid faster shortens or removes the ~$220K dip."),
 ("AFP 1 unpaid", "$29,039 expected in the current month.", "If it slips, October is that much lower."),
 ("30% plan sets not yet billed", "1.01.2 and 1.01.7 ($80,892) earned in Sep, assumed billed now and paid next month.", "Needs AFP 2 to go out this month."),
 ("Actual costs to date", "No payments recorded yet; June-Sept design and LNTP Annex A are treated as paid now.", "Recording actuals moves those payments into the months they really happened."),
 ("Subcontracts not yet awarded", "Construction placeholders: Net 30, 10% retainage, spread over the schedule.", "Real subs replace the placeholders with their own SOV and terms."),
 ("Switchgear equipment", "No cost loaded; 5.06 ($388K) bills on delivery.", "Loading the PO adds its cost and moves 50% of 5.06 to PO issue (Dec 2026)."),
 ("AHC Annex A staffing", "Runs to Nov 2027 per the AHC sheet; construction now finishes Aug 2027.", "Aligning staffing with the schedule would reduce late-2027 cost."),
 ("SOV 1.01.1 vs AFP 1", "AFP 1 billed $29,038.82 against a $29,038.72 line.", "10 cents; confirm which is right."),
]
for i, (item, now, eff) in enumerate(O, 1):
    r = 4 + i
    wo.cell(r, 1, i).font = f_n
    wo.cell(r, 2, item).font = f_b
    for c, t in ((3, now), (4, eff)):
        x = wo.cell(r, c, t); x.font = f_n; x.alignment = Alignment(wrap_text=True, vertical="top")
    wo.row_dimensions[r].height = 42

for s in wb.worksheets:
    s.sheet_view.showGridLines = False
    s.sheet_properties.pageSetUpPr.fitToPage = True
    s.page_setup.orientation = "landscape"
    s.page_setup.fitToWidth = 1
    s.page_setup.fitToHeight = 0
from openpyxl.workbook.properties import CalcProperties
wb.calculation = CalcProperties(fullCalcOnLoad=True)
wb.save(OUT)
print("saved", OUT, "months", len(months))
